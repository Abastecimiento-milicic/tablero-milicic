# ==============================================================================
# CONTROL DE PARTIDAS ABIERTAS SAP - CUENTA EM/RF (2101011001)
# Bases de Datos: PARTIDAS MLAR, MMAR, U003 y U365 (.xlsx)
# ==============================================================================

import streamlit as st
import pandas as pd
import numpy as np
import io
import datetime
import unicodedata
import os
import json
import plotly.express as px
import plotly.graph_objects as go
from openpyxl.worksheet.table import Table, TableStyleInfo

# ------------------------------------------------------------------------------
# CONFIGURACIÓN MANUAL: FECHA DE ACTUALIZACIÓN DEL REPORTE
# Modifica esta variable si deseas escribir la fecha directamente en el código:
# ------------------------------------------------------------------------------
FECHA_ACTUALIZACION_MANUAL = "05/10/2026"

# ------------------------------------------------------------------------------
# 1. CONFIGURACIÓN DE PÁGINA Y ESTILOS
# ------------------------------------------------------------------------------
st.set_page_config(
    page_title="Control EM/RF SAP - MLAR, MMAR, U003 y U365",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="auto"
)

# Estilos CSS personalizados
st.markdown("""
    <style>
    .main-title {
        font-size: 1.8rem;
        font-weight: 700;
        color: #1B365D;
        line-height: 1.35;
        margin-top: 0.5rem;
        margin-bottom: 0.2rem;
    }
    .sub-title {
        color: #475569;
        font-size: 0.95rem;
        margin-bottom: 1.2rem;
    }
    /* Contenedor principal: margen optimizado para aprovechar al máximo el ancho de la pantalla */
    .block-container {
        padding-top: 4rem !important;
        padding-bottom: 2rem !important;
        padding-left: 2.5rem !important;
        padding-right: 1.5rem !important;
        max-width: 100% !important;
    }
    /* ==========================================================================
       BARRA DE DESPLAZAMIENTO VERTICAL EN AZUL Y TABLA COMPLETA
       ========================================================================== */
    /* Permitir que la tabla sea visible y no recorte contenido ni la barra de herramientas */
    [data-testid="stDataFrame"] {
        overflow: visible !important;
        width: 100% !important;
    }

    /* Soporte estándar moderno (Chrome 121+, Edge, Firefox): Scrollbar vertical azul */
    * {
        scrollbar-color: #2563EB #EFF6FF !important;
    }

    [data-testid="stDataFrame"],
    [data-testid="stDataFrame"] *,
    [class*="dvn-"],
    [class*="dvn-"] * {
        scrollbar-color: #2563EB #EFF6FF !important;
    }

    /* Barra horizontal sutil solo si la pantalla es menor al ancho de la tabla */
    ::-webkit-scrollbar:horizontal,
    [data-testid="stDataFrame"] ::-webkit-scrollbar:horizontal,
    [data-testid="stDataFrame"] *::-webkit-scrollbar:horizontal,
    [class*="dvn-"]::-webkit-scrollbar:horizontal,
    [class*="dvn-"] *::-webkit-scrollbar:horizontal {
        height: 6px !important;
    }

    ::-webkit-scrollbar-thumb:horizontal,
    [data-testid="stDataFrame"] ::-webkit-scrollbar-thumb:horizontal,
    [data-testid="stDataFrame"] *::-webkit-scrollbar-thumb:horizontal,
    [class*="dvn-"]::-webkit-scrollbar-thumb:horizontal,
    [class*="dvn-"] *::-webkit-scrollbar-thumb:horizontal {
        background: #BFDBFE !important;
        border-radius: 4px !important;
    }

    /* Barra vertical azul en navegadores WebKit */
    ::-webkit-scrollbar {
        width: 12px !important;
        height: 0px !important;
    }

    ::-webkit-scrollbar:vertical {
        width: 12px !important;
    }

    ::-webkit-scrollbar-track:vertical {
        background: #EFF6FF !important;
        border-radius: 6px !important;
        border: 1px solid #BFDBFE !important;
    }

    ::-webkit-scrollbar-thumb:vertical {
        background: #2563EB !important;
        border-radius: 6px !important;
        border: 2px solid #EFF6FF !important;
    }

    ::-webkit-scrollbar-thumb:vertical:hover {
        background: #1D4ED8 !important;
    }

    [data-testid="stDataFrame"] ::-webkit-scrollbar:vertical,
    [data-testid="stDataFrame"] *::-webkit-scrollbar:vertical,
    [class*="dvn-"]::-webkit-scrollbar:vertical,
    [class*="dvn-"] *::-webkit-scrollbar:vertical {
        width: 12px !important;
    }

    [data-testid="stDataFrame"] ::-webkit-scrollbar-track:vertical,
    [data-testid="stDataFrame"] *::-webkit-scrollbar-track:vertical,
    [class*="dvn-"]::-webkit-scrollbar-track:vertical,
    [class*="dvn-"] *::-webkit-scrollbar-track:vertical {
        background: #EFF6FF !important;
        border-radius: 6px !important;
        border: 1px solid #BFDBFE !important;
    }

    [data-testid="stDataFrame"] ::-webkit-scrollbar-thumb:vertical,
    [data-testid="stDataFrame"] *::-webkit-scrollbar-thumb:vertical,
    [class*="dvn-"]::-webkit-scrollbar-thumb:vertical,
    [class*="dvn-"] *::-webkit-scrollbar-thumb:vertical {
        background: #2563EB !important;
        border-radius: 6px !important;
        border: 2px solid #EFF6FF !important;
    }

    [data-testid="stDataFrame"] ::-webkit-scrollbar-thumb:vertical:hover,
    [data-testid="stDataFrame"] *::-webkit-scrollbar-thumb:vertical:hover,
    [class*="dvn-"]::-webkit-scrollbar-thumb:vertical:hover,
    [class*="dvn-"] *::-webkit-scrollbar-thumb:vertical:hover {
        background: #1D4ED8 !important;
    }

    /* ==========================================================================
       BOTÓN DESCARGAR EXCEL LLAMATIVO EN AZUL
       ========================================================================== */
    div.stDownloadButton button {
        background: linear-gradient(135deg, #1D4ED8 0%, #2563EB 50%, #0284C7 100%) !important;
        color: #FFFFFF !important;
        font-weight: 700 !important;
        font-size: 1.05rem !important;
        border: none !important;
        border-radius: 8px !important;
        padding: 10px 24px !important;
        box-shadow: 0 4px 14px rgba(37, 99, 235, 0.45) !important;
        transition: all 0.25s ease-in-out !important;
    }

    div.stDownloadButton button:hover {
        background: linear-gradient(135deg, #1E40AF 0%, #1D4ED8 50%, #0369A1 100%) !important;
        color: #FFFFFF !important;
        box-shadow: 0 6px 20px rgba(37, 99, 235, 0.65) !important;
        transform: translateY(-2px) !important;
    }

    div.stDownloadButton button:active {
        transform: translateY(1px) !important;
        box-shadow: 0 2px 8px rgba(37, 99, 235, 0.4) !important;
    }

    div.stDownloadButton button p {
        color: #FFFFFF !important;
        font-weight: 700 !important;
    }
    </style>
""", unsafe_allow_html=True)

# ------------------------------------------------------------------------------
# 2. CARGA Y PREPROCESAMIENTO DE DATOS CON CACHÉ
# ------------------------------------------------------------------------------
def clean_col(c):
    nfkd = unicodedata.normalize('NFKD', str(c))
    return ''.join([ch for ch in nfkd if not unicodedata.combining(ch)]).lower().strip()

@st.cache_data(show_spinner=False)
def load_raw_data():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    sociedades = [
        ('MLAR', os.path.join(base_dir, 'PARTIDAS MLAR.xlsx')),
        ('MMAR', os.path.join(base_dir, 'PARTIDAS MMAR.xlsx')),
        ('U003', os.path.join(base_dir, 'PARTIDAS U003.xlsx')),
        ('U365', os.path.join(base_dir, 'PARTIDAS U365.xlsx'))
    ]
    
    op_map = {}
    map_path = os.path.join(base_dir, 'operadores_map.json')
    if os.path.exists(map_path):
        try:
            with open(map_path, 'r', encoding='utf-8') as f_op:
                op_map = json.load(f_op)
        except Exception:
            pass

    docs_list = []
    
    for soc, fpath in sociedades:
        if not os.path.exists(fpath):
            st.warning(f"Archivo {fpath} no encontrado en el directorio de trabajo.")
            continue
            
        raw = pd.read_excel(fpath, sheet_name='Data')
        raw = raw.rename(columns={c: clean_col(c) for c in raw.columns})
        
        # Filtrar filas de totales o vacías en columnas clave
        key_cols = ['documento compras', 'posicion', 'clave contabiliz.']
        raw = raw.dropna(subset=key_cols).copy()
        
        raw['Sociedad'] = soc
        
        # Mapear columnas estandarizadas
        raw['Pedido'] = raw['documento compras'].astype(np.int64).astype(str)
        raw['Posicion'] = raw['posicion'].astype(int).astype(str)
        raw['Key'] = soc + '_' + raw['Pedido'] + '_' + raw['Posicion']
        
        # Importe contable
        raw['Importe'] = pd.to_numeric(raw['importe en moneda local'], errors='coerce').fillna(0.0)
        raw['Moneda'] = raw.get('moneda local', pd.Series('ARS', index=raw.index)).fillna('ARS')
        
        # Claves y clases
        raw['Clave_Contab'] = raw['clave contabiliz.'].astype(int).astype(str)
        raw['Clase_Doc'] = raw.get('clase de documento', pd.Series('', index=raw.index)).fillna('').astype(str).str.strip()
        raw['Doc_SAP'] = raw.get('no documento', pd.Series('', index=raw.index)).fillna(0).astype(np.int64).astype(str)
        
        # Fechas
        raw['Fecha_Doc'] = pd.to_datetime(raw.get('fecha de documento'), errors='coerce', dayfirst=True)
        raw['Fe_Contabilizacion'] = pd.to_datetime(raw.get('fe.contabilizacion', raw.get('fecha de contabilizacion', pd.Series(pd.NaT, index=raw.index))), errors='coerce', dayfirst=True)
        raw['Clave_Referencia'] = raw.get('clave de referencia', raw.get('referencia', raw.get('asignacion', pd.Series('', index=raw.index)))).fillna('').astype(str).str.strip()
        raw['Fecha_Emision_OC'] = pd.to_datetime(raw.get('fecha emision oc'), errors='coerce', dayfirst=True)
        
        f_entrega = raw.get('fecha de entrega', raw.get('fecha entrega oc', raw.get('fecha de entrega oc', pd.Series(pd.NaT, index=raw.index))))
        raw['Fecha_Entrega_OC'] = pd.to_datetime(f_entrega, errors='coerce', dayfirst=True)
        raw['Fecha_Aprobacion_OC'] = pd.to_datetime(raw.get('fecha aprobacion final oc'), errors='coerce', dayfirst=True)
        raw['Primera_Fecha_Entrega_OC'] = pd.to_datetime(raw.get('primera fecha entrega oc'), errors='coerce', dayfirst=True)
        
        # Proveedor y Operadores
        prov_series = raw.get('nombre', raw.get('proveedor', raw.get('proveedor/centro suministrador', pd.Series('', index=raw.index))))
        raw['Proveedor'] = prov_series.fillna('DESCONOCIDO').astype(str).str.strip()
        
        # Operador OC: última columna de la base de datos ('OPERADOR DE OC', 'OPERADOR DE CO', 'OPERADOR OC')
        op_candidates = [
            'operador de oc', 'operador de co', 'operador oc', 'operador_oc',
            'operador_co', 'operador co', 'operador'
        ]
        col_operador_oc = None
        for cand in op_candidates:
            if cand in raw.columns:
                col_operador_oc = cand
                break
                
        # Si no se encontró por nombre explícito, usar la última columna de la base de datos
        if col_operador_oc is None and len(raw.columns) > 0:
            col_operador_oc = raw.columns[-1]

        if col_operador_oc is not None and col_operador_oc in raw.columns:
            raw['Operador_OC'] = raw[col_operador_oc].fillna(raw['Key'].map(op_map)).fillna('SIN ASIGNAR').astype(str).str.strip()
        else:
            raw['Operador_OC'] = raw['Key'].map(op_map).fillna('SIN ASIGNAR').astype(str).str.strip()
            
        raw['Operador_OC'] = raw['Operador_OC'].replace({'nan': 'SIN ASIGNAR', '': 'SIN ASIGNAR', 'None': 'SIN ASIGNAR'})
        mask_sin = raw['Operador_OC'] == 'SIN ASIGNAR'
        if mask_sin.any():
            fallback_map = raw.loc[mask_sin, 'Key'].map(op_map)
            raw.loc[mask_sin, 'Operador_OC'] = fallback_map.fillna('SIN ASIGNAR')
            
        raw['Operador_VA'] = raw.get('operador de va', raw.get('operador_va', pd.Series('SIN ASIGNAR', index=raw.index))).fillna('SIN ASIGNAR').astype(str).str.strip()
        
        # División y Grupo de Compras
        raw['Division'] = raw.get('division', pd.Series('S/D', index=raw.index)).fillna('S/D').astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
        gc_series = raw.get('grupo de compras', raw.get('grupo de compra oc', raw.get('grupo compras', pd.Series('SIN GRUPO', index=raw.index))))
        raw['Grupo_Compras'] = gc_series.fillna('SIN GRUPO').astype(str).str.strip()
        
        docs_list.append(raw)
        
    if not docs_list:
        return pd.DataFrame()
        
    full_docs = pd.concat(docs_list, ignore_index=True)
    return full_docs

def build_grouped_positions(df_docs, today_date):
    if df_docs.empty:
        return pd.DataFrame()
        
    def first_valid(series, default=''):
        for val in series:
            if pd.notna(val) and str(val).strip() not in ['', 'None', 'nan', 'DESCONOCIDO', 'SIN ASIGNAR', 'SIN GRUPO']:
                return str(val).strip()
        return default

    def join_divisions(series):
        divs = sorted(set(str(x).strip() for x in series if pd.notna(x) and str(x).strip() not in ['', 'S/D', 'nan']))
        return '/'.join(divs) if divs else 'S/D'

    grouped = df_docs.groupby(['Sociedad', 'Pedido', 'Posicion']).agg(
        Proveedor=('Proveedor', lambda s: first_valid(s, 'DESCONOCIDO')),
        Operador_OC=('Operador_OC', lambda s: first_valid(s, 'SIN ASIGNAR')),
        Operador_VA=('Operador_VA', lambda s: first_valid(s, 'SIN ASIGNAR')),
        Division=('Division', join_divisions),
        Grupo_Compras=('Grupo_Compras', lambda s: first_valid(s, 'SIN GRUPO')),
        Fe_Contabilizacion=('Fe_Contabilizacion', lambda s: s.dropna().max() if len(s.dropna()) > 0 else pd.NaT),
        Clave_Referencia=('Clave_Referencia', lambda s: first_valid(s, '')),
        Fecha_Emision_OC=('Fecha_Emision_OC', 'first'),
        Fecha_Entrega_OC=('Fecha_Entrega_OC', lambda s: s.dropna().iloc[0] if len(s.dropna()) > 0 else pd.NaT),
        Fecha_Aprobacion_OC=('Fecha_Aprobacion_OC', 'first'),
        Cant_Docs=('Clase_Doc', 'count'),
        Total_Debe=('Importe', lambda s: s[s > 0].sum()),
        Total_Haber=('Importe', lambda s: s[s < 0].sum()),
        Saldo_Neto=('Importe', 'sum'),
        Clases_Doc=('Clase_Doc', lambda s: ', '.join(sorted(set(str(x) for x in s if x)))),
        Fecha_Min_Doc=('Fecha_Doc', 'min'),
        Fecha_Max_Doc=('Fecha_Doc', 'max')
    ).reset_index()

    grouped['Total_Debe'] = grouped['Total_Debe'].round(2)
    grouped['Total_Haber'] = grouped['Total_Haber'].round(2)
    grouped['Saldo_Neto'] = grouped['Saldo_Neto'].round(2)
    grouped['Key'] = grouped['Sociedad'] + '_' + grouped['Pedido'] + '_' + grouped['Posicion']

    # Diagnósticos contables EM/RF
    def diag_emrf(r):
        saldo = r['Saldo_Neto']
        debe = r['Total_Debe']
        haber = r['Total_Haber']
        if abs(saldo) < 0.01:
            return 'COMPENSABLE (Saldo $0)'
        elif saldo < 0:
            if debe == 0:
                return 'FALTA FACTURA (Solo Recepción)'
            else:
                return 'FALTA FACTURA (Recepción parcial sin facturar)'
        else:
            if haber == 0:
                return 'FALTA RECEPCIÓN (Solo Factura)'
            else:
                return 'FALTA RECEPCIÓN (Factura mayor a Recepción)'

    def group_emrf(diag):
        if 'FALTA FACTURA' in diag:
            return 'TIENE RECEPCIÓN - FALTA FACTURA'
        elif 'FALTA RECEPCIÓN' in diag:
            return 'TIENE FACTURA - FALTA RECEPCIÓN'
        else:
            return 'COMPENSABLE (Saldo $0)'

    def venc_status(fec):
        if pd.isna(fec):
            return 'SIN FECHA OC'
        elif fec < today_date:
            return 'VENCIDA'
        else:
            return 'VIGENTE'

    def dias_atraso(fec):
        if pd.isna(fec):
            return 0
        delta = (today_date - fec).days
        return delta if delta > 0 else 0

    def calc_prioridad(r):
        emrf = r['Grupo_EMRF']
        venc = r['Estado_Vencimiento']
        if emrf == 'TIENE FACTURA - FALTA RECEPCIÓN':
            if venc == 'VENCIDA':
                return '🔴 URGENTE: OC Vencida sin Recepción'
            else:
                return '🟠 Factura sin Recepción (En plazo)'
        elif emrf == 'TIENE RECEPCIÓN - FALTA FACTURA':
            if venc == 'VENCIDA':
                return '🟡 RECLAMAR: OC Vencida sin Factura'
            else:
                return '🔵 Recepción en plazo (Pendiente Factura)'
        else:
            return '🟢 Listo para compensar (F.13)'

    grouped['Diagnostico_EMRF'] = grouped.apply(diag_emrf, axis=1)
    grouped['Grupo_EMRF'] = grouped['Diagnostico_EMRF'].apply(group_emrf)
    grouped['Estado_Vencimiento'] = grouped['Fecha_Entrega_OC'].apply(venc_status)
    grouped['Dias_Atraso'] = grouped['Fecha_Entrega_OC'].apply(dias_atraso)
    grouped['Prioridad_Accion'] = grouped.apply(calc_prioridad, axis=1)

    return grouped

# ==============================================================================
# GUÍA Y EXPLICACIÓN GENERAL (Comentado temporalmente por solicitud del usuario)
# Descomentar este bloque y las llamadas en sidebar/encabezado cuando se desee reactivar.
# ==============================================================================
# def render_general_explanation():
#     st.markdown("### 📘 Guía y Explicación General")
#     st.markdown("""
#     **¿Qué analiza este tablero?**  
#     Supervisa la cuenta puente de SAP **EM/RF (`2101011001`)**, conciliando lo ingresado físicamente en almacén/planta (**MIGO**) contra lo facturado por el proveedor (**MIRO**).
# 
#     ---
#     #### 📌 ¿Qué significan los pedidos según su estado?
#     * 📦 **TIENE RECEPCIÓN - FALTA FACTURA:**  
#       El almacén ya recepcionó la mercadería o servicio (*Recepcionado Haber* con saldo negativo), pero el proveedor aún no facturó o la factura no fue cargada (*Facturado Debe* en `$0,00`).  
#       👉 **Acción:** Reclamar la factura al proveedor o gestionar la contabilización en Cuentas por Pagar.
# 
#     * 📄 **TIENE FACTURA - FALTA RECEPCIÓN:**  
#       La factura del proveedor ya fue contabilizada en SAP (*Facturado Debe* positivo), pero en el almacén no consta el ingreso de mercadería o remito (*Recepcionado Haber* en `$0,00`).  
#       👉 **Acción:** Reclamar el remito o ingreso físico al almacén/planta para evitar riesgo de pagar sin recibir.
# 
#     * 🟢 **COMPENSABLE (Saldo $0):**  
#       Lo facturado coincide exactamente con lo recepcionado (Saldo Neto `$0,00`).  
#       👉 **Acción:** Listo para depurar y cerrar en SAP mediante la transacción **`F.13`**.
# 
#     ---
#     #### ⚠️ Aclaración sobre la columna «ESTADO»
#     La columna **ESTADO** de la tabla indica el cumplimiento de la **fecha de entrega pactada** (`FE OC`):
#     * **VENCIDA:** La fecha prometida de entrega del pedido ya expiró respecto a la fecha actual.
#     * **VIGENTE:** El pedido aún se encuentra en plazo para su entrega.  
#     *(Nota: Esta columna mide el vencimiento de la fecha de entrega del pedido, no si le falta factura o recepción; eso se consulta en el filtro de Estado Contable EM/RF).*
# 
#     ---
#     #### ⚙️ Filtros iniciales (al abrir el sistema)
#     * Por defecto, la aplicación muestra únicamente pedidos con **«TIENE RECEPCIÓN - FALTA FACTURA»** correspondientes a los Grupos de Compras **C01 y C07**.
#     * Podés cambiar el estado o quitar los grupos de compras desde el **Panel de Control (menú lateral izquierdo)** para ver el universo completo o ver los que les falta recepción.
#     """)

# Cargar datos
with st.spinner("Cargando bases de datos de SAP (MLAR, MMAR, U003, U365)... "):
    df_docs_all = load_raw_data()

if df_docs_all.empty:
    st.error("No se encontraron datos en los archivos de bases de datos SAP (MLAR, MMAR, U003, U365).")
    st.stop()

# ------------------------------------------------------------------------------
# 3. BARRA LATERAL (SIDEBAR) - FILTROS OPERATIVOS
# ------------------------------------------------------------------------------
st.sidebar.title("Panel de Control")

# [EXPLICACIÓN GENERAL COMENTADA TEMPORALMENTE - Descomentar para reactivar]
# with st.sidebar.popover("ℹ️ Explicación General", help="Haz clic para ver la guía y cómo interpretar los pedidos"):
#     render_general_explanation()

# Fecha de actualización escrita a mano (editable en pantalla o en la constante de código)
fecha_actualizacion = st.sidebar.text_input(
    "📅 Fecha de Actualización:",
    value=FECHA_ACTUALIZACION_MANUAL,
    help="Ingresa la fecha de corte/actualización a mano (formato DD/MM/AAAA o texto)."
).strip()

# Filtro 1: Selección de Sociedad / Base de Datos
soc_disponibles = sorted(df_docs_all['Sociedad'].unique())
soc_options = ["Todas (Consolidado)"] + soc_disponibles
sel_soc = st.sidebar.selectbox("Base de Datos / Sociedad:", soc_options, index=0)

if sel_soc == "Todas (Consolidado)":
    df_docs = df_docs_all.copy()
else:
    df_docs = df_docs_all[df_docs_all['Sociedad'] == sel_soc].copy()

# Fecha de corte basada en la fecha ingresada a mano (para cálculo de vencimientos y días de atraso)
try:
    corte_timestamp = pd.to_datetime(fecha_actualizacion, dayfirst=True)
    if pd.isna(corte_timestamp):
        corte_timestamp = pd.Timestamp.now().normalize()
except Exception:
    corte_timestamp = pd.Timestamp.now().normalize()

# Construir agrupaciones basadas en la fecha de corte
df_pos = build_grouped_positions(df_docs, corte_timestamp)

st.sidebar.markdown("---")
st.sidebar.subheader("Filtros de Búsqueda")

# Buscador unificado: Pedido o Proveedor
search_query = st.sidebar.text_input("Buscar Pedido o Proveedor:", placeholder="Ej: 4500121598 o AURELIA").strip()

# Filtro 2: Estado Contable EM/RF
emrf_all = ["TIENE RECEPCIÓN - FALTA FACTURA", "TIENE FACTURA - FALTA RECEPCIÓN", "COMPENSABLE (Saldo $0)"]
sel_emrf = st.sidebar.multiselect("Estado Contable EM/RF:", emrf_all, default=["TIENE RECEPCIÓN - FALTA FACTURA"])

# Filtro 3: Grupo de Compras
grupos_disp = sorted([g for g in df_pos['Grupo_Compras'].unique() if g != 'SIN GRUPO'])
default_grupos = [g for g in ["C01", "C07"] if g in grupos_disp]
sel_grupos = st.sidebar.multiselect("Grupo de Compras OC:", grupos_disp, default=default_grupos)

# Filtro 4: Estado de Vencimiento
venc_all = ["VENCIDA", "VIGENTE", "SIN FECHA OC"]
sel_venc = st.sidebar.multiselect("Vencimiento de OC:", venc_all, default=[])

# Filtro 5: Proveedor
proveedores_disponibles = sorted([p for p in df_pos['Proveedor'].unique() if p != 'DESCONOCIDO'])
sel_prov = st.sidebar.multiselect("Proveedor:", proveedores_disponibles, default=[])

# Filtro 6: Operador de Compras OC
operadores_oc = sorted([o for o in df_pos['Operador_OC'].unique() if o != 'SIN ASIGNAR'])
sel_oper_oc = st.sidebar.multiselect("Operador Compras (OC):", operadores_oc, default=[])

# Filtro 7: Operador de VA (Verificación Facturas)
operadores_va = sorted([o for o in df_pos['Operador_VA'].unique() if o != 'SIN ASIGNAR'])
sel_oper_va = st.sidebar.multiselect("Operador Verif. Factura (VA):", operadores_va, default=[])

# Filtro 8: Slider de Días de Atraso
max_dias = int(df_pos['Dias_Atraso'].max()) if not df_pos.empty else 0
sel_dias_min = st.sidebar.slider("Días de Atraso mínimos:", min_value=0, max_value=max(max_dias, 1), value=0, step=15)

# APLICAR FILTROS
filtered = df_pos.copy()

if search_query:
    q = search_query.lower()
    filtered = filtered[
        filtered['Pedido'].str.lower().str.contains(q, na=False) |
        filtered['Proveedor'].str.lower().str.contains(q, na=False)
    ]

if sel_emrf:
    filtered = filtered[filtered['Grupo_EMRF'].isin(sel_emrf)]

if sel_venc:
    filtered = filtered[filtered['Estado_Vencimiento'].isin(sel_venc)]

if sel_prov:
    filtered = filtered[filtered['Proveedor'].isin(sel_prov)]

if sel_oper_oc:
    filtered = filtered[filtered['Operador_OC'].isin(sel_oper_oc)]

if sel_oper_va:
    filtered = filtered[filtered['Operador_VA'].isin(sel_oper_va)]

if sel_grupos:
    filtered = filtered[filtered['Grupo_Compras'].isin(sel_grupos)]

if sel_dias_min > 0:
    filtered = filtered[filtered['Dias_Atraso'] >= sel_dias_min]

# Sidebar summary
st.sidebar.markdown("---")
pct_filtrado = (len(filtered) / len(df_pos)) * 100 if len(df_pos) > 0 else 0
st.sidebar.markdown(f"**Posiciones seleccionadas:** {len(filtered):,d} / {len(df_pos):,d} (`{pct_filtrado:.1f}%`)")
st.sidebar.markdown(f"**Saldo Neto Total:** `${filtered['Saldo_Neto'].sum():,.2f}` ARS")

# ------------------------------------------------------------------------------
# 4. ENCABEZADO
# ------------------------------------------------------------------------------
soc_badge = f"<span style='background:#1B365D; color:white; padding:4px 10px; border-radius:12px; font-size:0.85rem;'>{sel_soc}</span>"
st.markdown(f"<div class='main-title'>Control de Partidas Abiertas SAP - Cuenta EM/RF {soc_badge}</div>", unsafe_allow_html=True)
st.markdown(f"<div class='sub-title'>Cuenta de Compensación <b>2101011001</b> | Fecha de Actualización: <b>{fecha_actualizacion}</b> | Fuentes: <b>MLAR, MMAR, U003 y U365</b></div>", unsafe_allow_html=True)

# [EXPLICACIÓN GENERAL COMENTADA TEMPORALMENTE - Descomentar para reactivar el botón de ayuda en el encabezado]
# with st.popover("ℹ️ Explicación General", help="Haz clic para ver la guía y cómo interpretar los pedidos"):
#     render_general_explanation()

# ------------------------------------------------------------------------------
# 5. TABLA DE TODAS LAS POSICIONES Y DESCARGA
# ------------------------------------------------------------------------------
cols_display = [
    'Sociedad',
    'Pedido',
    'Posicion',
    'Proveedor',
    'Grupo_Compras',
    'Operador_OC',
    'Fe_Contabilizacion',
    'Dias_Atraso',
    'Total_Debe',
    'Total_Haber',
    'Saldo_Neto',
    'Fecha_Entrega_OC',
    'Estado_Vencimiento'
]

rename_cols_excel = {
    'Sociedad': 'Sociedad',
    'Pedido': 'Pedido',
    'Posicion': 'POS',
    'Proveedor': 'Proveedor',
    'Grupo_Compras': 'GC',
    'Operador_OC': 'Operador_OC',
    'Fe_Contabilizacion': 'FE CONT.',
    'Dias_Atraso': 'Dias_Atraso',
    'Total_Debe': 'Facturado (Debe)',
    'Total_Haber': 'Recepcionado (Haber)',
    'Saldo_Neto': 'Saldo_Neto',
    'Fecha_Entrega_OC': 'FE OC',
    'Estado_Vencimiento': 'ESTADO'
}

# Preparar archivo Excel para descarga inmediata
output = io.BytesIO()
df_excel = filtered[cols_display].rename(columns=rename_cols_excel).copy()
with pd.ExcelWriter(output, engine='openpyxl') as writer:
    df_excel.to_excel(writer, index=False, sheet_name='Posiciones_Filtradas')
    ws = writer.sheets['Posiciones_Filtradas']
    
    # 1. Formatear fechas como DD/MM/YYYY
    date_cols = ['FE CONT.', 'FE OC']
    for col_name in date_cols:
        if col_name in df_excel.columns:
            col_idx = list(df_excel.columns).index(col_name) + 1
            for row_idx in range(2, len(df_excel) + 2):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    cell.number_format = 'DD/MM/YYYY'
                    
    # 2. Formatear valores monetarios
    curr_cols = ['Facturado (Debe)', 'Recepcionado (Haber)', 'Saldo_Neto']
    for col_name in curr_cols:
        if col_name in df_excel.columns:
            col_idx = list(df_excel.columns).index(col_name) + 1
            for row_idx in range(2, len(df_excel) + 2):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None and isinstance(cell.value, (int, float)):
                    cell.number_format = '$ #,##0.00'

    # 3. Formato oficial de Tabla de Excel con estilo azul y filtros
    if len(df_excel) > 0:
        tab = Table(displayName="PartidasAbiertas", ref=ws.dimensions)
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)

    # 4. Autoajustar el ancho de cada columna para que se lea todo el texto sin recortar
    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = 0
        for cell in col:
            if cell.value is not None:
                val = str(cell.value)
                if cell.number_format == 'DD/MM/YYYY':
                    val = 'DD/MM/YYYY'
                elif '$' in str(cell.number_format):
                    try:
                        val = f'$ {float(cell.value):,.2f}'
                    except Exception:
                        pass
                max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

excel_bytes = output.getvalue()
timestamp_str = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')

c_title, c_rows = st.columns([3, 1.2])
with c_title:
    st.markdown(f"### 📑 Todas las Posiciones ({len(filtered):,d})")
with c_rows:
    st.download_button(
        label="📥 Descargar Reporte en Excel",
        data=excel_bytes,
        file_name=f"CONTROL_EM_RF_{sel_soc}_{timestamp_str}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width='stretch',
        type='primary'
    )
    filas_opciones = [10, 25, 50, 100, 250, "Todas"]
    cant_filas = st.selectbox("Filas en pantalla:", filas_opciones, index=filas_opciones.index("Todas"))

# Configuración visual para st.dataframe
col_config_dict = {
    'Sociedad': st.column_config.TextColumn("Sociedad", width=65),
    'Pedido': st.column_config.TextColumn("Pedido", width=90),
    'Posicion': st.column_config.TextColumn("POS", width=45),
    'Proveedor': st.column_config.TextColumn("Proveedor", width=145),
    'Grupo_Compras': st.column_config.TextColumn("GC", width=45),
    'Operador_OC': st.column_config.TextColumn("Operador_OC", width=95),
    'Fe_Contabilizacion': st.column_config.DateColumn("FE CONT.", format="DD/MM/YYYY", width=85),
    'Dias_Atraso': st.column_config.NumberColumn("Dias_Atraso", format="%d", width=80),
    'Total_Debe': st.column_config.NumberColumn("Facturado (Debe)", format="$ %.2f", alignment="left", width=110),
    'Total_Haber': st.column_config.NumberColumn("Recepcionado (Haber)", format="$ %.2f", alignment="left", width=120),
    'Saldo_Neto': st.column_config.NumberColumn("Saldo_Neto", format="$ %.2f", alignment="left", width=100),
    'Fecha_Entrega_OC': st.column_config.DateColumn("FE OC", format="DD/MM/YYYY", width=80),
    'Estado_Vencimiento': st.column_config.TextColumn("ESTADO", width=80)
}

if cant_filas == "Todas":
    df_mostrar = filtered[cols_display]
else:
    df_mostrar = filtered[cols_display].head(int(cant_filas))

st.dataframe(
    df_mostrar,
    column_config=col_config_dict,
    width='stretch',
    hide_index=True
)

if cant_filas == "Todas":
    st.caption(f"Mostrando el 100% de las posiciones filtradas ({len(df_mostrar):,d} registros).")
else:
    st.caption(f"Mostrando las primeras {len(df_mostrar):,d} de {len(filtered):,d} posiciones filtradas. (Selecciona **'Todas'** en el menú superior para ver y recorrer todas las posiciones según el filtro).")
